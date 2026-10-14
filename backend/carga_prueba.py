"""
carga_prueba.py
----------------
Prueba de carga: genera 500 documentos sintéticos (PDF, Excel, Word e
imágenes), los sube por la API real del sistema y mide si las respuestas
se vuelven más lentas al crecer el índice.

Todos los archivos se llaman PRUEBA_CARGA_xxxx_... y se suben con la
categoría "otro" (así no aparecen en las secciones públicas del inicio).
El texto usa un vocabulario distinto al de los datos reales (equipos,
inventarios, mantenimiento) para no contaminar las respuestas del asistente.

Uso (con el backend corriendo en http://localhost:8000, desde backend/):
    python carga_prueba.py generar               # crea los 500 archivos en carga_prueba/docs
    python carga_prueba.py medir --etiqueta antes  # mide velocidad ANTES de subir
    python carga_prueba.py subir                 # sube los 500 y mide cada subida
    python carga_prueba.py medir --etiqueta despues
    python carga_prueba.py informe               # compara antes vs después
    python carga_prueba.py limpiar               # borra TODOS los PRUEBA_CARGA_* del sistema
"""

import argparse
import io
import json
import os
import random
import re
import statistics
import textwrap
import time
from pathlib import Path

import requests
from docx import Document as DocxDocument
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from services.auth_service import create_access_token

# 127.0.0.1 y no "localhost": en Windows `requests` prueba primero IPv6 (::1) y espera
# ~2 s a que falle antes de usar IPv4, lo que suma 2 s falsos a CADA llamada.
BASE_URL = os.getenv("BASE_URL", "http://127.0.0.1:8000")
CARPETA = Path(__file__).parent / "carga_prueba"
DOCS = CARPETA / "docs"
PREFIJO = "PRUEBA_CARGA_"
TOTAL = 500  # 125 de cada tipo

# Palabras que NO pueden aparecer en el texto sintético: son las que usan las
# preguntas reales (o disparan la búsqueda de optativas, que busca "OPT").
PROHIBIDAS = re.compile(
    r"opt|billar|torneo|docente|horario|materia|semestre|base de datos|beca|convocatoria|programaci|sigla",
    re.I,
)

# ---------------------------------------------------------------- contenido

SUJETOS = [
    "El área de mantenimiento", "La comisión de infraestructura", "El comité de seguridad",
    "La unidad administrativa", "El almacén central", "La biblioteca especializada",
    "El taller de electrónica", "La sala de cómputo", "El departamento de compras",
    "La comisión de inventarios",
]
ACCIONES = [
    "realizó la revisión trimestral de", "registró el ingreso de", "solicitó la reposición de",
    "concluyó la verificación de", "programó el traslado de", "informó el deterioro de",
    "autorizó la baja de", "actualizó el registro de",
]
OBJETOS = [
    "los proyectores del bloque {b}", "las sillas del aula {n}", "los cables de red del piso {p}",
    "las pizarras acrílicas", "los extintores del pasillo {p}", "las baterías de respaldo",
    "los estabilizadores de voltaje", "los ventiladores de techo", "las mesas del laboratorio {n}",
    "los armarios metálicos", "los teclados y ratones", "las luminarias del corredor",
]
CIERRES = [
    "Se recomienda continuar con el seguimiento mensual.",
    "El informe fue remitido a la dirección para su conocimiento.",
    "La observación quedó registrada en el libro de novedades.",
    "Se adjunta el detalle en el anexo correspondiente.",
    "No se reportaron incidentes adicionales.",
    "La próxima revisión está prevista para el mes siguiente.",
]
EQUIPOS = [
    "Osciloscopio digital", "Fuente de poder regulable", "Multímetro de banco", "Router de borde",
    "Switch de 24 puertos", "Estación de soldadura", "Proyector multimedia", "Servidor en rack",
    "Impresora de red", "Cámara de vigilancia", "Batería de respaldo", "Monitor de 24 pulgadas",
    "Placa de desarrollo", "Analizador lógico", "Generador de señales",
]
ESTADOS = ["Operativo", "Operativo", "Operativo", "En reparación", "Dado de baja", "En calibración"]
RESPONSABLES = [
    "Quiroga Betanzos", "Ledezma Arancibia", "Zambrana Ortuño", "Fuentes Caballero",
    "Baldivieso Aliaga", "Terceros Maidana", "Saavedra Peñaranda", "Camacho Urquidi",
]
COLORES_FONDO = [(255, 250, 235), (235, 245, 255), (240, 255, 240), (255, 240, 245), (245, 245, 245)]


def _sentencia(rng: random.Random) -> str:
    objeto = rng.choice(OBJETOS).format(b=rng.choice("ABCDE"), n=rng.randint(1, 30), p=rng.randint(1, 4))
    return f"{rng.choice(SUJETOS)} {rng.choice(ACCIONES)} {objeto}. {rng.choice(CIERRES)}"


def _parrafos(rng: random.Random, cantidad: int) -> list[str]:
    return [" ".join(_sentencia(rng) for _ in range(rng.randint(2, 4))) for _ in range(cantidad)]


# ---------------------------------------------------------------- generadores

def hacer_pdf(titulo: str, parrafos: list[str]) -> bytes:
    """PDF mínimo escrito a mano (Helvetica, texto real seleccionable)."""
    lineas = [titulo, ""]
    for p in parrafos:
        lineas += textwrap.wrap(p, 92) + [""]
    paginas = [lineas[i:i + 46] for i in range(0, len(lineas), 46)]

    def esc(t: str) -> bytes:
        return t.encode("latin-1", "replace").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")

    objetos: list[bytes] = []  # el índice 0 es el objeto 1
    n_pag = len(paginas)
    ids_paginas = [4 + 2 * i for i in range(n_pag)]
    objetos.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{i} 0 R" for i in ids_paginas)
    objetos.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pag} >>".encode())
    objetos.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, lin in enumerate(paginas):
        contenido = b"BT /F1 11 Tf 50 800 Td 16 TL\n" + b"".join(b"(" + esc(l) + b") Tj T*\n" for l in lin) + b"ET"
        objetos.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {5 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        objetos.append(f"<< /Length {len(contenido)} >>\nstream\n".encode() + contenido + b"\nendstream")

    salida = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objetos, start=1):
        offsets.append(len(salida))
        salida += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    inicio_xref = len(salida)
    salida += f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        salida += f"{off:010d} 00000 n \n".encode()
    salida += f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{inicio_xref}\n%%EOF".encode()
    return bytes(salida)


def hacer_docx(titulo: str, parrafos: list[str], rng: random.Random) -> bytes:
    doc = DocxDocument()
    doc.add_heading(titulo, level=1)
    for p in parrafos:
        doc.add_paragraph(p)
    tabla = doc.add_table(rows=1, cols=3)
    tabla.rows[0].cells[0].text, tabla.rows[0].cells[1].text, tabla.rows[0].cells[2].text = "Equipo", "Cantidad", "Estado"
    for _ in range(5):
        c = tabla.add_row().cells
        c[0].text, c[1].text, c[2].text = rng.choice(EQUIPOS), str(rng.randint(1, 20)), rng.choice(ESTADOS)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def hacer_xlsx(rng: random.Random) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Codigo", "Equipo", "Laboratorio", "Estado", "Responsable", "Cantidad", "Fecha de revision"])
    for i in range(rng.randint(25, 60)):
        ws.append([
            f"INV-{rng.randint(10000, 99999)}", rng.choice(EQUIPOS), f"Sala {rng.randint(1, 30)}",
            rng.choice(ESTADOS), rng.choice(RESPONSABLES), rng.randint(1, 12),
            f"{rng.randint(1, 28):02d}/{rng.randint(1, 12):02d}/2026",
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fuente(tamano: int):
    for ruta in ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/segoeui.ttf"):
        try:
            return ImageFont.truetype(ruta, tamano)
        except OSError:
            continue
    return ImageFont.load_default()


def hacer_imagen(titulo: str, rng: random.Random) -> bytes:
    """Cartel con texto grande y nítido, como un aviso escaneado."""
    img = Image.new("RGB", (1000, 720), rng.choice(COLORES_FONDO))
    d = ImageDraw.Draw(img)
    d.text((50, 40), titulo, fill=(20, 20, 60), font=_fuente(42))
    y = 130
    for _ in range(6):
        for linea in textwrap.wrap(_sentencia(rng), 52)[:2]:
            d.text((50, y), linea, fill=(10, 10, 10), font=_fuente(28))
            y += 40
        y += 12
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def generar() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    rng = random.Random(2026)
    por_tipo = TOTAL // 4
    creados = 0
    t0 = time.time()
    for i in range(1, TOTAL + 1):
        tipo = ("pdf", "xlsx", "docx", "png")[i % 4]
        tema = rng.choice(["informe_mantenimiento", "acta_inventario", "reporte_equipos", "revision_almacen"])
        nombre = f"{PREFIJO}{i:04d}_{tema}.{tipo}"
        titulo = f"{tema.replace('_', ' ').title()} N {i:04d}/2026"
        if tipo == "pdf":
            datos = hacer_pdf(titulo, _parrafos(rng, rng.randint(4, 8)))
            texto_control = titulo
        elif tipo == "docx":
            datos = hacer_docx(titulo, _parrafos(rng, rng.randint(4, 8)), rng)
            texto_control = titulo
        elif tipo == "xlsx":
            datos = hacer_xlsx(rng)
            texto_control = ""
        else:
            datos = hacer_imagen(titulo, rng)
            texto_control = titulo
        assert not PROHIBIDAS.search(texto_control), nombre
        (DOCS / nombre).write_bytes(datos)
        creados += 1
    print(f"Se generaron {creados} archivos ({por_tipo} de cada tipo) en {DOCS} ({time.time() - t0:.0f} s)")


# ---------------------------------------------------------------- API

def _headers() -> dict:
    token = create_access_token({"sub": os.getenv("ADMIN_EMAIL"), "role": "admin"})
    return {"Authorization": f"Bearer {token}"}


MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "png": "image/png",
}


def _todos_los_documentos() -> list[dict]:
    """GET /admin/documents está paginado: se recorren todas las páginas (de 100 en 100)."""
    docs, pagina = [], 1
    while True:
        r = requests.get(f"{BASE_URL}/admin/documents", params={"page": pagina, "page_size": 100}, headers=_headers()).json()
        docs += r["items"]
        if len(docs) >= r["total"] or not r["items"]:
            return docs
        pagina += 1


def subir() -> None:
    archivos = sorted(DOCS.glob(f"{PREFIJO}*"))
    if not archivos:
        raise SystemExit("No hay archivos: corré primero `python carga_prueba.py generar`")
    ya = {d["filename"] for d in _todos_los_documentos()}
    pendientes = [a for a in archivos if a.name not in ya]
    print(f"{len(pendientes)} por subir ({len(archivos) - len(pendientes)} ya estaban subidos)")

    tiempos = []  # (nombre, tipo, segundos, chunks)
    errores = []
    t_total = time.time()
    for n, ruta in enumerate(pendientes, start=1):
        ext = ruta.suffix[1:]
        t = time.time()
        r = requests.post(
            f"{BASE_URL}/admin/documents",
            headers=_headers(),
            files={"file": (ruta.name, ruta.read_bytes(), MIME[ext])},
            data={"category": "otro"},
            timeout=300,
        )
        seg = time.time() - t
        if r.status_code == 200:
            tiempos.append((ruta.name, ext, seg, r.json()["chunk_count"]))
        else:
            errores.append((ruta.name, r.status_code, r.text[:200]))
        if n % 25 == 0:
            print(f"  {n}/{len(pendientes)} subidos — promedio últimos 25: {statistics.mean(x[2] for x in tiempos[-25:]):.2f} s", flush=True)

    total = time.time() - t_total
    resumen = {"subidos": len(tiempos), "errores": errores, "segundos_totales": total, "detalle": tiempos}
    (CARPETA / "resultados_subida.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\nSubida terminada: {len(tiempos)} OK, {len(errores)} con error, {total / 60:.1f} min en total")
    for tipo in ("pdf", "docx", "xlsx", "png"):
        xs = [x[2] for x in tiempos if x[1] == tipo]
        if xs:
            print(f"  {tipo:5s} n={len(xs):3d}  promedio {statistics.mean(xs):.2f} s  máx {max(xs):.2f} s  fragmentos={sum(x[3] for x in tiempos if x[1] == tipo)}")
    if len(tiempos) >= 200:
        primeros, ultimos = [x[2] for x in tiempos[:100]], [x[2] for x in tiempos[-100:]]
        print(f"  primeras 100 subidas: {statistics.mean(primeros):.2f} s | últimas 100: {statistics.mean(ultimos):.2f} s")
    for e in errores[:10]:
        print("  ERROR", e)


# Preguntas de la medición. Cada una usa un camino distinto de búsqueda.
PREGUNTAS = [
    ("filtro exacto", "que materias hay en primer semestre"),
    ("filtro exacto", "que materias hay en septimo semestre"),
    ("texto (docente)", "que docente dicta base de datos"),
    ("texto (imagen OCR)", "cuando es el torneo de billar"),
    ("optativas", "cuales son las materias optativas"),
    ("semantica", "como puedo postular a la beca de alimentacion"),
]


def _chat(pregunta: str) -> tuple[float, int, bool, int]:
    """Devuelve (segundos, status, desde_cache, reintentos). Si Groq limita por tokens/minuto, espera y reintenta."""
    reintentos = 0
    while True:
        t = time.time()
        r = requests.post(f"{BASE_URL}/chat", json={"pregunta": pregunta}, headers=_headers(), timeout=120)
        seg = time.time() - t
        if r.status_code == 200 or reintentos >= 3:
            return seg, r.status_code, bool(r.json().get("desde_cache")) if r.status_code == 200 else False, reintentos
        reintentos += 1
        time.sleep(25)  # límite de tokens por minuto de Groq


def _cronometrar(metodo: str, ruta: str, repeticiones: int = 5, **kw) -> float:
    xs = []
    for _ in range(repeticiones):
        t = time.time()
        requests.request(metodo, f"{BASE_URL}{ruta}", headers=_headers(), timeout=120, **kw).raise_for_status()
        xs.append(time.time() - t)
    return statistics.median(xs)


def _vaciar_cache_via_api() -> None:
    """Subir y borrar un documento mínimo hace que el backend vacíe el caché de
    respuestas, así la medición empieza con todas las preguntas "nuevas"."""
    d = DocxDocument()
    d.add_paragraph("documento temporal para vaciar el cache")
    buf = io.BytesIO()
    d.save(buf)
    r = requests.post(f"{BASE_URL}/admin/documents", headers=_headers(), data={"category": "otro"},
                      files={"file": (f"{PREFIJO}TEMP.docx", buf.getvalue(), MIME["docx"])}, timeout=120)
    r.raise_for_status()
    requests.delete(f"{BASE_URL}/admin/documents/{r.json()['id']}", headers=_headers(), timeout=120).raise_for_status()


def medir(etiqueta: str) -> None:
    _vaciar_cache_via_api()
    stats = requests.get(f"{BASE_URL}/admin/stats", headers=_headers()).json()
    print(f"[{etiqueta}] documentos={stats['total_documentos']}  fragmentos={stats['total_fragmentos']}")
    res = {"documentos": stats["total_documentos"], "fragmentos": stats["total_fragmentos"], "chat": [], "api": {}}

    print("\nChat (respuesta nueva; se espera entre preguntas por el límite de tokens/min de Groq):")
    for camino, pregunta in PREGUNTAS:
        seg, status, cache, reint = _chat(pregunta)
        res["chat"].append({"camino": camino, "pregunta": pregunta, "seg": seg, "status": status, "cache": cache, "reintentos": reint})
        print(f"  {seg:5.1f} s  [{status}] {camino:20s} {pregunta}" + (f"  (reintentos por límite Groq: {reint})" if reint else ""))
        time.sleep(12)

    print("\nMisma pregunta otra vez (respuesta desde el caché que aprende):")
    seg, status, cache, _ = _chat(PREGUNTAS[0][1])
    res["cache_seg"] = seg
    print(f"  {seg:5.2f} s  cache={cache}")

    print("\nEndpoints (mediana de 5):")
    for nombre, metodo, ruta in [
        ("listar documentos (admin)", "GET", "/admin/documents"),
        ("estadísticas", "GET", "/admin/stats"),
        ("usuarios pag. 1", "GET", "/admin/users?page=1&page_size=10"),
        ("documentos públicos: comunicados", "GET", "/documentos?seccion=comunicados"),
    ]:
        res["api"][nombre] = _cronometrar(metodo, ruta)
        print(f"  {res['api'][nombre] * 1000:7.0f} ms  {nombre}")

    CARPETA.mkdir(exist_ok=True)
    (CARPETA / f"resultados_{etiqueta}.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")


def informe() -> None:
    a = json.loads((CARPETA / "resultados_antes.json").read_text(encoding="utf-8"))
    d = json.loads((CARPETA / "resultados_despues.json").read_text(encoding="utf-8"))
    print(f"Documentos: {a['documentos']} -> {d['documentos']}   Fragmentos: {a['fragmentos']} -> {d['fragmentos']}\n")
    print(f"{'Consulta':52s} {'antes':>8s} {'después':>8s}")
    for x, y in zip(a["chat"], d["chat"]):
        print(f"{x['camino'] + ': ' + x['pregunta'][:36]:52s} {x['seg']:7.1f}s {y['seg']:7.1f}s")
    print(f"{'Pregunta repetida (caché)':52s} {a['cache_seg']:7.2f}s {d['cache_seg']:7.2f}s")
    for k in a["api"]:
        print(f"{k:52s} {a['api'][k] * 1000:6.0f}ms {d['api'][k] * 1000:6.0f}ms")


def limpiar() -> None:
    docs = [x for x in _todos_los_documentos() if x["filename"].startswith(PREFIJO)]
    print(f"Borrando {len(docs)} documentos de prueba…")
    for n, doc in enumerate(docs, start=1):
        requests.delete(f"{BASE_URL}/admin/documents/{doc['id']}", headers=_headers(), timeout=120).raise_for_status()
        if n % 50 == 0:
            print(f"  {n}/{len(docs)}", flush=True)
    print("Listo. (Los archivos generados siguen en carga_prueba/docs; podés borrar esa carpeta.)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("accion", choices=["generar", "subir", "medir", "informe", "limpiar"])
    ap.add_argument("--etiqueta", default="antes")
    args = ap.parse_args()
    {"generar": generar, "subir": subir, "informe": informe, "limpiar": limpiar}.get(args.accion, lambda: medir(args.etiqueta))()
